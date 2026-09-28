"""Generate the starter packing-list Word template (and the field reference docs).

    python tools/make_template.py          # -> assets/templates/packing_list_template.docx
    python tools/make_template.py --docs   # also refresh docs/TEMPLATE_FIELDS.md and the README.txt section

The template is committed; this script only exists so it can be reproduced.
Users edit their own copy in Word (templates/packing_list_template.docx).
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from tallycraft.packing_docx import (PARTS_TABLE_ALT_TEXT, field_reference_markdown,  # noqa: E402
                                     field_reference_text)

OUT = ROOT / "assets" / "templates" / "packing_list_template.docx"
FONT = "Calibri"
INK = RGBColor(0x1F, 0x1F, 0x1F)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
RED = RGBColor(0x8A, 0x00, 0x00)
LINE = "A6A6A6"
SHADE = "EDEDED"
CONTENT_W = 7.5  # inches: US Letter minus 0.5 in margins


# ---------------------------------------------------------------- low-level helpers

def run(p, text, size=10, bold=False, color=INK, font=FONT):
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    return r


def para(container, text="", size=10, bold=False, color=INK, align=None, after=0, before=0, font=FONT):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_after, pf.space_before, pf.line_spacing = Pt(after), Pt(before), 1.0
    if align:
        p.alignment = align
    if text:
        run(p, text, size, bold, color, font)
    return p


def label(container, text):
    return para(container, text.upper(), size=7.5, bold=True, color=GREY, after=1)


def tag(container, text):
    """A docxtpl paragraph tag ({%p ... %}); it must be alone in its paragraph."""
    return para(container, text, size=8, color=GREY)


def first_para(cell):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    return p


def set_cell_shading(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def set_cell_margins(table, top=60, bottom=60, left=90, right=90):
    pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    pr.append(mar)


def set_borders(table, outer=None, inner=None):
    """outer/inner: (size in 1/8 pt, hex color) or None for no line."""
    pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        spec = outer if side in ("top", "left", "bottom", "right") else inner
        el = OxmlElement(f"w:{side}")
        if spec:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec[0]))
            el.set(qn("w:color"), spec[1])
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    pr.append(borders)


def fixed_table(doc, rows, widths):
    t = doc.add_table(rows=rows, cols=len(widths))
    t.autofit = False
    pr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    pr.append(layout)
    for row in t.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w)
    for col, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(round(w * 1440)))
    return t


def row_flag(row, name):
    """w:tblHeader (repeat on each page) or w:cantSplit (keep a row on one page)."""
    pr = row._tr.get_or_add_trPr()
    el = OxmlElement(f"w:{name}")
    pr.append(el)


def para_box(p, border="A6A6A6", fill=None, size=6):
    pr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), border)
        bdr.append(el)
    pr.append(bdr)
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        pr.append(shd)


def field(p, instr, size=8, color=GREY):
    """A complex Word field (PAGE / NUMPAGES). Every run carries the same formatting,
    so Word keeps the size and color when it updates the number."""
    def fmt_run():
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), FONT)
        fonts.set(qn("w:hAnsi"), FONT)
        rpr.append(fonts)
        col = OxmlElement("w:color")
        col.set(qn("w:val"), str(color))
        rpr.append(col)
        for tag_name in ("w:sz", "w:szCs"):
            sz = OxmlElement(tag_name)
            sz.set(qn("w:val"), str(int(size * 2)))
            rpr.append(sz)
        r.append(rpr)
        return r

    for kind in ("begin", "instr", "separate", "result", "end"):
        r = fmt_run()
        if kind in ("begin", "separate", "end"):
            ch = OxmlElement("w:fldChar")
            ch.set(qn("w:fldCharType"), kind)
            r.append(ch)
        elif kind == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instr} "
            r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = "1"
            r.append(t)
        p._p.append(r)


def cell_text(cell, text, size=9, bold=False, align=None, color=INK):
    p = first_para(cell)
    if align:
        p.alignment = align
    if text:
        run(p, text, size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return p


# ---------------------------------------------------------------- the template

def build() -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(0)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.5))
    sec.footer_distance = Inches(0.3)

    # Footer on every page: order reference left, "Page X of Y" right.
    fp = sec.footer.paragraphs[0]
    stops = fp.paragraph_format.tab_stops
    for pos in (3.25, 6.5):  # the built-in Footer style's center/right tabs would catch the tab first
        stops.add_tab_stop(Inches(pos), WD_TAB_ALIGNMENT.CLEAR)
    stops.add_tab_stop(Inches(CONTENT_W), WD_TAB_ALIGNMENT.RIGHT)
    run(fp, "{{ footer_reference }}\tPage ", 8, color=GREY)
    field(fp, "PAGE", color=GREY)
    run(fp, " of ", 8, color=GREY)
    field(fp, "NUMPAGES", color=GREY)

    body = doc  # add to the document body
    body.paragraphs[0]._p.getparent().remove(body.paragraphs[0]._p) if body.paragraphs else None

    # ---- header: logo + shop name | PACKING LIST + date
    head = fixed_table(body, 1, [4.6, 2.9])
    set_borders(head)
    set_cell_margins(head, 0, 0, 0, 0)
    left, right = head.rows[0].cells
    first_para(left)
    left.paragraphs[0]._p.getparent().remove(left.paragraphs[0]._p)
    tag(left, "{%p if shop_logo %}")
    para(left, "{{ shop_logo }}", after=4)
    tag(left, "{%p endif %}")
    para(left, "{{ shop_name }}", size=16, bold=True)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    cell_text(right, "PACKING LIST", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    para(right, "Date: {{ packing_date }}", size=9, color=GREY, align=WD_ALIGN_PARAGRAPH.RIGHT)
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    para(body, after=8)

    # ---- customer block
    cust = fixed_table(body, 1, [4.6, 2.9])
    set_borders(cust, outer=(6, LINE))
    set_cell_margins(cust, 110, 130, 160, 160)
    ship, order = cust.rows[0].cells
    p = first_para(ship)
    run(p, "SHIP TO", 7.5, True, GREY)
    para(ship, "{{ customer_name }}", size=12, bold=True, before=1)
    tag(ship, "{%p if customer_address %}")
    para(ship, "{{ customer_address }}", size=10, before=1)
    tag(ship, "{%p endif %}")
    first_para(order)
    order.paragraphs[0]._p.getparent().remove(order.paragraphs[0]._p)
    tag(order, "{%p if etsy_order_number %}")
    label(order, "Etsy order #")
    para(order, "{{ etsy_order_number }}", size=12, bold=True, after=6)
    tag(order, "{%p endif %}")
    tag(order, "{%p if order_date %}")
    label(order, "Order date")
    para(order, "{{ order_date }}", size=10)
    tag(order, "{%p endif %}")
    para(order)  # keeps the cell valid for Word even when both fields are empty
    para(body, after=6)

    p = para(body, after=8)
    run(p, "Items ordered: ", 10, True)
    run(p, "{{ items_ordered }}", 10)

    # ---- total weight (prominent)
    tw = fixed_table(body, 1, [2.6, 4.9])
    set_borders(tw, outer=(16, "1F1F1F"))
    set_cell_margins(tw, 120, 140, 160, 160)
    a, b = tw.rows[0].cells
    cell_text(a, "TOTAL SHIPPING WEIGHT", size=8, bold=True, color=GREY)
    cell_text(b, "{{ total_weight_lb_oz }}   ({{ total_weight_g }} g)", size=20, bold=True,
              align=WD_ALIGN_PARAGRAPH.RIGHT)
    tag(body, "{%p if unmeasured_warning %}")
    w = para(body, "{{ unmeasured_warning }}", size=9.5, bold=True, color=RED, before=6)
    para_box(w, border="8A0000", fill="FDECEC", size=8)
    tag(body, "{%p endif %}")
    para(body, after=8)

    tag(body, "{%p if kit_legend %}")
    para(body, "{{ kit_legend }}", size=7.5, color=GREY, after=4)
    tag(body, "{%p endif %}")

    # ---- parts table. Columns 2 and 4 hold the {%tc %} tags (removed when filled in);
    # column 3 repeats once per kit. Alt Text lets TallyCraft restore these widths.
    widths = [0.95, 2.85, 0.05, 0.85, 0.05, 0.6, 0.75, 0.8, 0.45]
    t = fixed_table(body, 5, widths)
    set_borders(t, outer=(6, LINE), inner=(4, LINE))
    set_cell_margins(t, 50, 50, 70, 70)
    desc = OxmlElement("w:tblDescription")
    desc.set(qn("w:val"), PARTS_TABLE_ALT_TEXT)
    t._tbl.tblPr.append(desc)
    C, R = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT
    heads = ["Picture", "Part", "{%tc for h in kit_headers %}", "{{ h }}", "{%tc endfor %}", "Total count",
             "Weight per piece (g)", "Total weight (g)", "Done"]
    for cell, text in zip(t.rows[0].cells, heads):
        cell_text(cell, text, size=8, bold=True, align=C)
        set_cell_shading(cell, SHADE)
    row_flag(t.rows[0], "tblHeader")
    cell_text(t.rows[1].cells[0], "{%tr for p in parts %}", size=8, color=GREY)
    body_cells = t.rows[2].cells
    cell_text(body_cells[0], "{{ p.picture }}", align=C)
    para(body_cells[0], "{{ p.dimensions }}", size=6.5, color=GREY, align=C)
    cell_text(body_cells[1], "{{ p.name }}{{ p.marker }}", size=9)
    cell_text(body_cells[2], "{%tc for c in p.counts %}", size=8, color=GREY)
    cell_text(body_cells[3], "{{ c }}", size=10, align=C)
    cell_text(body_cells[4], "{%tc endfor %}", size=8, color=GREY)
    cell_text(body_cells[5], "{{ p.total_count }}", size=10, align=C)
    cell_text(body_cells[6], "{{ p.weight_per_piece }}", size=10, align=R)
    cell_text(body_cells[7], "{{ p.total_weight }}", size=10, align=R)
    box = cell_text(body_cells[8], "", align=C)
    run(box, "☐", 16, font="Segoe UI Symbol")
    row_flag(t.rows[2], "cantSplit")
    cell_text(t.rows[3].cells[0], "{%tr endfor %}", size=8, color=GREY)
    tot = ["", "{{ total_label }}", "{%tc for k in kit_totals %}", "{{ k }}", "{%tc endfor %}",
           "{{ total_count }}", "", "{{ total_weight_measured }}", ""]
    aligns = [None, None, None, C, None, C, None, R, None]
    for cell, text, al in zip(t.rows[4].cells, tot, aligns):
        cell_text(cell, text, size=9.5, bold=True, align=al)
        set_cell_shading(cell, SHADE)
    row_flag(t.rows[4], "cantSplit")

    tag(body, "{%p for f in footnotes %}")
    para(body, "{{ f.marker }} {{ f.name }}: could not be measured, so its weight isn't in the total. "
               "Reason: {{ f.reason }}", size=7.5, color=RED, before=2)
    tag(body, "{%p endfor %}")

    # ---- sign-off, note, shop record
    para(body, "Packed by: ______________________________      Date: ____________________", size=10, before=20)
    tag(body, "{%p if note_to_customer %}")
    n = para(body, "{{ note_to_customer }}", size=10, before=14)
    para_box(n, border=LINE, size=6)
    tag(body, "{%p endif %}")
    para(body, "{{ shop_record }}", size=7, color=GREY, before=18)
    return doc


def write_docs() -> None:
    (ROOT / "docs" / "TEMPLATE_FIELDS.md").write_text(field_reference_markdown() + "\n", encoding="utf-8",
                                                      newline="\n")
    readme = ROOT / "dist_template" / "README.txt"
    text = readme.read_text(encoding="utf-8")
    start, end = "<<TEMPLATE FIELDS START>>", "<<TEMPLATE FIELDS END>>"
    if start in text and end in text:
        a, b = text.index(start) + len(start), text.index(end)
        text = text[:a] + "\n" + field_reference_text() + "\n" + text[b:]
        readme.write_text(text, encoding="utf-8", newline="\n")


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    build().save(OUT)
    print(f"Wrote {OUT}")
    if "--docs" in sys.argv[1:]:
        write_docs()
        print("Refreshed docs/TEMPLATE_FIELDS.md and the README.txt field section")


if __name__ == "__main__":
    main()
