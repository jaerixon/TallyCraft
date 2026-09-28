"""Fill an editable Word template (docxtpl / Jinja2) from a packing-list record.

The record (packing.build_record) is the source of truth; this module turns it
into template fields, part pictures (PNG, high resolution), and a filled-in
.docx. PDF conversion is separate (word_pdf.py) because it needs Microsoft Word.

Every placeholder is defined once in FIELDS. That single list drives template
validation, docs/TEMPLATE_FIELDS.md, the README.txt section, and the app's
Template Help window. See SPEC §9.
"""

from __future__ import annotations

import math
import os
import re
from dataclasses import dataclass
from pathlib import Path

from .calc import format_lb_oz
from .packing import calibration_text, decode_logo, part_name, stamp

PARTS_TABLE_ALT_TEXT = "TallyCraft parts table"
TOTALS_TABLE_ALT_TEXT = "TallyCraft totals row"  # optional: the totals as a one-row table inside the parts table
PICTURE_DPI = 400  # printed resolution of part pictures (the brief asks for at least 300)
PICTURE_HEIGHT_IN = 0.5
SUPERSAMPLE = 3  # draw big, then shrink: smooth lines without an anti-aliasing library


def bundled_template() -> Path:
    """The starter template shipped with the app (inside the exe when packaged)."""
    import sys
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[2]))
    return base / "assets" / "templates" / "packing_list_template.docx"


# ============================================================================ the field list

@dataclass(frozen=True)
class Field:
    name: str
    description: str
    example: str
    group: str = "Order and customer"


FIELDS: list[Field] = [
    # shop
    Field("shop_name", "Your shop name (File > Settings). Empty if not set.", "Bun Homes Co.", "Shop"),
    Field("shop_logo", "Your logo image (File > Settings), 0.8 in tall. Empty if not set. Put it on its own line.",
          "(image)", "Shop"),
    Field("packing_date", "The date the packing list was created.", "2026-09-28", "Shop"),
    # customer / order
    Field("customer_name", "Customer name (always filled in).", "Jordan Rivera"),
    Field("customer_address", "Shipping address; its line breaks are kept. Empty if not entered.",
          "123 Maple Street / Springfield, IL 62704"),
    Field("etsy_order_number", "Etsy order number. Empty if not entered.", "3141592653"),
    Field("order_date", "Order date. Empty if not entered.", "2026-09-27"),
    Field("items_ordered", "The packages in the order with quantities.", "2x XL Standard Box, 1x XL Tunnel + Ramp"),
    Field("note_to_customer", "The note to the customer. Empty if none.", "Thank you for your order!"),
    Field("footer_reference", "For the page footer: \"Etsy order #…\", or \"Packing list — <customer>\" if there's "
                              "no order number.", "Etsy order #3141592653"),
    # weight
    Field("total_weight_lb_oz", "Total shipping weight in pounds and ounces (measured parts).", "4 lb 14.9 oz",
          "Weight"),
    Field("total_weight_g", "Total shipping weight in grams, 1 decimal, no unit.", "2,237.2", "Weight"),
    Field("unmeasured_count", "How many parts couldn't be measured (0 if none).", "1", "Weight"),
    Field("unmeasured_warning", "A warning sentence when parts couldn't be measured; empty otherwise. Use with "
                                "{%p if unmeasured_warning %} … {%p endif %}.",
          "WARNING: Total weight excludes 1 part that could not be measured (marked * below). It still goes in "
          "the box.", "Weight"),
    # parts table
    Field("parts", "The parts table rows. Loop over them with {%tr for p in parts %} … {%tr endfor %}; each p has "
                   "the p.… fields below.", "(list)", "Parts table"),
    Field("kit_headers", "Kit column headings, for the column loop {%tc for h in kit_headers %} {{ h }} "
                         "{%tc endfor %}. Full kit names, \"Kit 1…\", or one per package, depending on what fits.",
          "XL Standard Box (1 of 2), XL Standard Box (2 of 2)", "Parts table"),
    Field("kit_legend", "Explains short kit headings (\"Kit 1 = …\"); empty when full names fit.",
          "Kit columns: Kit 1 = XL Standard Box (1 of 4); …", "Parts table"),
    Field("kit_totals", "Totals-row values for the kit columns (same column loop).", "28, 28, 9", "Parts table"),
    Field("total_count", "Totals row: all parts across every kit.", "65", "Parts table"),
    Field("total_weight_measured", "Totals row: total weight in grams, 2 decimals (measured parts).", "12,564.18",
          "Parts table"),
    Field("total_label", "Totals row label: \"Total\", or \"Total (measured parts)\" when some couldn't be "
                         "measured.", "Total", "Parts table"),
    Field("footnotes", "One per unmeasured part, for {%p for f in footnotes %} … {%p endfor %}; each f has "
                       "f.marker, f.name, f.reason.", "(list)", "Parts table"),
    # shop record
    Field("calibration_method", "\"Reference piece\" or \"Control sample\".", "Reference piece", "Shop record"),
    Field("calibration_details", "Full calibration sentence (method, values, ratio, preset).",
          "Calibration: control sample 4 × 2 in, weight 23.456 g → 0.45446 g/cm² (2.93200 g/in²).", "Shop record"),
    Field("calibration_preset", "Calibration preset name, if the values still match it. Empty otherwise.",
          "DixiePly 120526", "Shop record"),
    Field("grams_per_cm2", "Weight per area, g/cm², 5 decimals.", "0.45446", "Shop record"),
    Field("grams_per_in2", "Weight per area, g/in², 5 decimals.", "2.93200", "Shop record"),
    Field("calculated_at", "When Calculate was pressed.", "2026-09-28 10:15", "Shop record"),
    Field("created_at", "When the packing list was created.", "2026-09-28 10:20", "Shop record"),
    Field("app_version", "TallyCraft version.", "0.5.0", "Shop record"),
    Field("shop_record", "The whole small-print shop record in one line.",
          "Calibration: … Calculated 2026-09-28 10:15. Packing list created 2026-09-28 10:20. TallyCraft 0.5.0.",
          "Shop record"),
]

PART_FIELDS: list[Field] = [
    Field("picture", "Drawing of the part (cut lines solid, engrave lines light), fitted to the Picture column.",
          "(image)", "Each part (p.…)"),
    Field("dimensions", "Bounding box size, in the unit chosen in Settings (\"Show dimensions in\").",
          "14.500 × 14.250 in", "Each part (p.…)"),
    Field("display_name", "The name to show: the Display name set in the pieces table, or the file name "
                          "(without .dxf) if none is set.", "Front Panel", "Each part (p.…)"),
    Field("name", "The file name without .dxf (always, even if a Display name is set).",
          "XL v1 - Main Box - Front Panel v2 - x3", "Each part (p.…)"),
    Field("marker", "Footnote marker for unmeasured parts (\" *1\"); empty otherwise.", " *1", "Each part (p.…)"),
    Field("counts", "This part's count in each kit column, for {%tc for c in p.counts %} {{ c }} {%tc endfor %}; "
                    "\"—\" if a kit doesn't use it.", "2, 2, —", "Each part (p.…)"),
    Field("total_count", "Count across every kit in the order.", "4", "Each part (p.…)"),
    Field("weight_per_piece", "Grams per piece, 2 decimals; \"—\" if unmeasured.", "423.64", "Each part (p.…)"),
    Field("total_weight", "Grams for all of this part, 2 decimals; \"—\" if unmeasured.", "1,694.58",
          "Each part (p.…)"),
    Field("skipped", "True if the part couldn't be measured.", "False", "Each part (p.…)"),
    Field("reason", "Why it couldn't be measured; empty otherwise.", "The outline has a gap…", "Each part (p.…)"),
]

FOOTNOTE_FIELDS: list[Field] = [
    Field("marker", "The footnote marker.", "*1", "Each footnote (f.…)"),
    Field("name", "Part name (Display name if set, otherwise the file name).", "Broken gap",
          "Each footnote (f.…)"),
    Field("reason", "Why it couldn't be measured.", "The outline has a gap…", "Each footnote (f.…)"),
]

KNOWN_FIELDS = {f.name for f in FIELDS}
PART_ATTRS = {f.name for f in PART_FIELDS}
FOOTNOTE_ATTRS = {f.name for f in FOOTNOTE_FIELDS}
REQUIRED_HELP = ("The parts table loop ({%tr for p in parts %} … {%tr endfor %}) and the total weight "
                 "({{ total_weight_lb_oz }} or {{ total_weight_g }}) must be in the template.")


def field_reference_markdown() -> str:
    """docs/TEMPLATE_FIELDS.md — generated from FIELDS (tests check it's up to date)."""
    out = ["# Packing list template fields", "",
           "<!-- Generated from src/tallycraft/packing_docx.py (FIELDS). Run `python tools/make_template.py "
           "--docs` after changing fields. -->", "",
           "TallyCraft fills a Word template (`templates/packing_list_template.docx`) using "
           "[docxtpl](https://docxtpl.readthedocs.io/) placeholders. Edit the template in Word; keep the "
           "placeholders (including their `{{ }}` / `{% %}` brackets) intact.", "",
           "**Required:** " + REQUIRED_HELP, "",
           "**Tags:** `{{ name }}` inserts a value. `{%p if … %}` / `{%p endif %}` show or hide whole paragraphs. "
           "`{%tr for … %}` / `{%tr endfor %}` repeat table rows. `{%tc for … %}` / `{%tc endfor %}` repeat table "
           "columns. Each `{%tr %}` / `{%tc %}` / `{%p %}` tag must sit alone in its own row / cell / paragraph.",
           "",
           f"**Column widths:** keep the parts table's Alt Text set to `{PARTS_TABLE_ALT_TEXT}` (Word: right-click "
           "the table > Table Properties > Alt Text). TallyCraft then keeps your column widths and shares the "
           "remaining width among the kit columns.", ""]
    for group in dict.fromkeys(f.group for f in FIELDS + PART_FIELDS + FOOTNOTE_FIELDS):
        out += [f"## {group}", "", "| Placeholder | Contains | Example |", "|---|---|---|"]
        for f in FIELDS + PART_FIELDS + FOOTNOTE_FIELDS:
            if f.group != group:
                continue
            prefix = "p." if f in PART_FIELDS else "f." if f in FOOTNOTE_FIELDS else ""
            desc = f.description.replace("|", "\\|")
            out.append(f"| `{{{{ {prefix}{f.name} }}}}` | {desc} | {f.example.replace('|', '/')} |")
        out.append("")
    return "\n".join(out)


def field_reference_text(width: int = 76) -> str:
    """Plain-text version for README.txt and the Template Help window."""
    import textwrap
    lines = ["Placeholders are written like {{ customer_name }} in the Word template.",
             "Required: " + REQUIRED_HELP, "",
             "Tags: {{ name }} inserts a value. {%p if ... %} / {%p endif %} show or hide",
             "paragraphs. {%tr for ... %} / {%tr endfor %} repeat table rows, and",
             "{%tc for ... %} / {%tc endfor %} repeat table columns. Each {%tr %}, {%tc %}",
             "and {%p %} tag must sit alone in its own row, cell, or paragraph.",
             f"Keep the parts table's Alt Text \"{PARTS_TABLE_ALT_TEXT}\" so TallyCraft",
             "can keep your column widths.", ""]
    for group in dict.fromkeys(f.group for f in FIELDS + PART_FIELDS + FOOTNOTE_FIELDS):
        lines += [group.upper(), "-" * len(group)]
        for f in FIELDS + PART_FIELDS + FOOTNOTE_FIELDS:
            if f.group != group:
                continue
            prefix = "p." if f in PART_FIELDS else "f." if f in FOOTNOTE_FIELDS else ""
            lines.append(f"  {{{{ {prefix}{f.name} }}}}")
            lines += textwrap.wrap(f.description, width, initial_indent="      ", subsequent_indent="      ")
            lines.append(f"      e.g. {f.example}")
        lines.append("")
    return "\n".join(lines)


# ============================================================================ pictures

def render_picture_png(picture: dict, path: Path, width_in: float, height_in: float = PICTURE_HEIGHT_IN,
                       dpi: int = PICTURE_DPI) -> None:
    """Draw a part picture: cut outline solid black, holes thinner, engrave light
    grey, unclosed runs dashed. Scaled to fit its own box."""
    from PIL import Image, ImageDraw

    W, H = max(1, round(width_in * dpi)), max(1, round(height_in * dpi))
    S = SUPERSAMPLE
    img = Image.new("RGB", (W * S, H * S), "white")
    draw = ImageDraw.Draw(img)
    loops = picture.get("loops") or []
    opens = picture.get("open") or []
    engrave = picture.get("engrave") or []
    pts = [p for l in loops for p in l] + [p for c in opens for p in c] + [p for e in engrave for p in e]
    if pts:
        minx, maxx = min(p[0] for p in pts), max(p[0] for p in pts)
        miny, maxy = min(p[1] for p in pts), max(p[1] for p in pts)
        pad = 0.04 * dpi * S
        bw, bh = max(maxx - minx, 1e-9), max(maxy - miny, 1e-9)
        scale = min((W * S - 2 * pad) / bw, (H * S - 2 * pad) / bh)
        ox = (W * S - bw * scale) / 2
        oy = (H * S - bh * scale) / 2

        def tx(p):  # image y grows downward; DXF y grows upward
            return ox + (p[0] - minx) * scale, oy + (maxy - p[1]) * scale

        def pt(points_width):  # line width in points -> pixels at this resolution
            return max(1, round(points_width / 72 * dpi * S))

        for line in engrave:
            if len(line) >= 2:
                draw.line([tx(p) for p in line], fill=(160, 160, 160), width=pt(0.35), joint="curve")
        depths = picture.get("depths") or [0] * len(loops)
        for loop, depth in sorted(zip(loops, depths), key=lambda ld: -ld[1]):
            if len(loop) >= 2:
                draw.line([tx(p) for p in loop] + [tx(loop[0])], fill=(0, 0, 0),
                          width=pt(0.9 if depth % 2 == 0 else 0.6), joint="curve")
        dash = 0.03 * dpi * S
        for chain in opens:
            for a, b in zip(chain, chain[1:]):
                (x1, y1), (x2, y2) = tx(a), tx(b)
                length = math.hypot(x2 - x1, y2 - y1)
                steps = max(1, int(length // dash))
                for i in range(0, steps, 2):
                    t0, t1 = i / steps, min(1.0, (i + 1) / steps)
                    draw.line([(x1 + (x2 - x1) * t0, y1 + (y2 - y1) * t0),
                               (x1 + (x2 - x1) * t1, y1 + (y2 - y1) * t1)], fill=(60, 60, 60), width=pt(0.6))
    img = img.resize((W, H), Image.LANCZOS)
    img.save(path, dpi=(dpi, dpi))


# ============================================================================ template geometry & kit columns

@dataclass
class TableGeometry:
    """Column widths (inches) of the template's parts table, split around the kit loop."""
    content_width: float
    before: list[float]  # columns before {%tc for %} (e.g. Picture, Part)
    kit: float  # the template's kit column width
    after: list[float]  # columns after {%tc endfor %} (Total count, weights, checkbox)
    name_index: int  # which "before" column is the flexible part-name column
    found: bool = True


DEFAULT_GEOMETRY = TableGeometry(7.5, [0.95, 2.0], 0.85, [0.6, 0.75, 0.8, 0.45], name_index=1, found=False)


def _cell_text(tc) -> str:
    from docx.oxml.ns import qn
    return "".join(t.text or "" for t in tc.iter(qn("w:t")))


def read_table_geometry(template_path) -> TableGeometry:
    """Find the parts table (by Alt Text) and read its column widths."""
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(str(template_path))
    sec = doc.sections[0]
    content = (sec.page_width - sec.left_margin - sec.right_margin) / 914400
    tbl = _find_parts_table(doc)
    if tbl is None:
        return TableGeometry(content, DEFAULT_GEOMETRY.before, DEFAULT_GEOMETRY.kit, DEFAULT_GEOMETRY.after, 1,
                             found=False)
    widths = [int(g.get(qn("w:w"))) / 1440 for g in tbl._tbl.tblGrid.findall(qn("w:gridCol"))]
    header = [_cell_text(tc) for tc in tbl.rows[0]._tr.findall(qn("w:tc"))]
    try:
        start = next(i for i, t in enumerate(header) if "{%tc for" in t.replace("  ", " "))
        end = next(i for i, t in enumerate(header) if "{%tc endfor" in t.replace("  ", " "))
    except StopIteration:
        return TableGeometry(content, widths, 0.0, [], name_index=min(1, len(widths) - 1))
    before, after = widths[:start], widths[end + 1:]
    kit = sum(widths[start + 1:end]) or DEFAULT_GEOMETRY.kit
    name_index = max(range(len(before)), key=lambda i: before[i]) if before else 0
    return TableGeometry(content, before, kit, after, name_index)


def _find_parts_table(doc, alt_text: str = PARTS_TABLE_ALT_TEXT):
    """The table (w:tbl element wrapper) with this Alt Text, including tables inside table cells."""
    from docx.oxml.ns import qn
    from docx.table import Table
    for tbl in doc.element.body.iter(qn("w:tbl")):
        pr = tbl.find(qn("w:tblPr"))
        desc = pr.find(qn("w:tblDescription")) if pr is not None else None
        if desc is not None and desc.get(qn("w:val")) == alt_text:
            return Table(tbl, doc)
    return None


def _measure(text: str, size_pt: float) -> float:
    """Width of `text` in points for the bold header font (Calibri, else Arial, else an estimate)."""
    from PIL import ImageFont
    windir = os.environ.get("WINDIR", r"C:\Windows")
    for name in ("calibrib.ttf", "arialbd.ttf"):
        path = os.path.join(windir, "Fonts", name)
        if os.path.exists(path):
            try:
                font = ImageFont.truetype(path, 100)
                return font.getlength(text) * size_pt / 100
            except OSError:
                pass
    return len(text) * size_pt * 0.55


def plan_kit_columns(kit_labels: list[str], packages: list[dict] | None, geo: TableGeometry,
                     header_pt: float = 8.0, name_min_in: float = 1.2) -> dict:
    """How kit columns fit on a portrait page, using the template's real widths.

    Returns {"mode": "full" | "short" | "per_package" | "omitted", "headers", "legend",
    "indices" (kit_counts index per column), "kit_width", "name_width"} (widths in inches).
    """
    fixed = sum(geo.before) - geo.before[geo.name_index] + sum(geo.after) if geo.before else sum(geo.after)
    room = geo.content_width - fixed

    def fit(labels):
        n = len(labels)
        if n == 0:
            return {"mode": "full", "headers": [], "kit_width": 0.0, "name_width": room}
        w = min(geo.kit, (room - name_min_in) / n)
        inner = (w - 0.1) * 72

        def ok(label):
            words = all(_measure(word, header_pt) <= inner for word in label.split())
            lines = math.ceil(_measure(label, header_pt) / max(inner, 1) * 1.15)
            return words and lines <= 3

        if w >= 0.55 and all(ok(l) for l in labels):
            return {"mode": "full", "headers": list(labels), "kit_width": w, "name_width": room - w * n}
        w = min(0.55, (room - 1.0) / n)
        if w >= 0.3:
            return {"mode": "short", "headers": [f"Kit {i}" for i in range(1, n + 1)], "kit_width": w,
                    "name_width": room - w * n}
        return None

    n = len(kit_labels)
    plan = fit(kit_labels)
    if plan:
        plan["legend"] = list(zip(plan["headers"], kit_labels)) if plan["mode"] == "short" else []
        plan["indices"] = list(range(n))
        return plan
    if packages and sum(int(p.get("quantity", 1)) for p in packages) == n:
        labels, indices, i = [], [], 0
        for p in packages:
            q = int(p.get("quantity", 1))
            labels.append(p["name"] if q == 1 else f"{p['name']} (each of {q})")
            indices.append(i)
            i += q
        plan = fit(labels)
        if plan:
            plan["legend"] = list(zip(plan["headers"], labels)) if plan["mode"] == "short" else []
            plan["mode"] = "per_package"
            plan["indices"] = indices
            return plan
    return {"mode": "omitted", "headers": [], "kit_width": 0.0, "name_width": room, "legend": [], "indices": []}


# ============================================================================ context

def _g(v, decimals=2) -> str:
    return "—" if v is None else f"{v:,.{decimals}f}"


DIMENSION_UNITS = {"in": "Inches", "mm": "Millimeters", "file": "Each file's own units"}
MM_PER = {"in": 25.4, "mm": 1.0}


def _dims(line: dict, show: str = "in") -> str:
    """Bounding box size under the picture. show: "in" / "mm" (convert), or "file" (as the file is)."""
    bb = (line.get("picture") or {}).get("bbox")
    if not bb:
        return ""
    unit = line.get("unit")
    if unit not in MM_PER:  # unknown units: can't convert, show the file's numbers
        return f"{bb[0]:.3f} × {bb[1]:.3f}"
    target = unit if show not in MM_PER else show
    w, h = (v * MM_PER[unit] / MM_PER[target] for v in bb)
    return f"{w:.3f} × {h:.3f} in" if target == "in" else f"{w:.1f} × {h:.1f} mm"


def build_context(record: dict, plan: dict, picture=lambda line: "", logo="", dim_units: str = "in") -> dict:
    """Template fields from a record. `picture(line)` / `logo` supply images (docxtpl
    InlineImage objects when rendering; plain strings in tests)."""
    cust, res, cal = record["customer"], record["results"], record["calibration"]
    order_no = str(cust.get("etsy_order") or "").strip()
    total = res.get("total_g") or 0.0
    skipped_n = int(res.get("skipped_count") or 0)
    parts, footnotes = [], []
    kit_totals = [0] * len(plan["indices"])
    for line in res["lines"]:
        all_counts = list(line.get("kit_counts") or [])
        counts = [all_counts[i] if i < len(all_counts) else None for i in plan["indices"]]
        for k, c in enumerate(counts):
            kit_totals[k] += c or 0
        marker = ""
        shown = (line.get("display_name") or "").strip() or part_name(line["name"])
        if line.get("skipped"):
            footnotes.append({"marker": f"*{len(footnotes) + 1}", "name": shown, "reason": line.get("reason", "")})
            marker = f" *{len(footnotes)}"
        parts.append({"picture": picture(line), "dimensions": _dims(line, dim_units),
                      "display_name": shown, "name": part_name(line["name"]),
                      "marker": marker, "counts": ["—" if c is None else str(c) for c in counts],
                      "total_count": str(line.get("total_count", "")),
                      "weight_per_piece": _g(line.get("weight_per_piece_g")),
                      "total_weight": _g(line.get("item_total_g")),
                      "skipped": bool(line.get("skipped")), "reason": line.get("reason", "") or ""})
    g = cal.get("grams_per_cm2")
    legend = "; ".join(f"{h} = {full}" for h, full in plan["legend"])
    if plan["mode"] == "short":
        legend = "Kit columns: " + legend
    elif plan["mode"] == "per_package":
        legend = (f"This order has {len(res.get('kit_labels') or [])} kits, too many for one column each, so "
                  "counts are shown once per package (every kit of a package has the same parts)."
                  + (" Columns: " + legend if legend else ""))
    elif plan["mode"] == "omitted":
        legend = (f"This order has {len(res.get('kit_labels') or [])} kits, too many to show count columns on a "
                  "portrait page, so only the Total count is shown.")
    they = "They still go" if skipped_n != 1 else "It still goes"
    details = calibration_text(cal)
    record_line = (f"{details} Calculated {stamp(record.get('calculated_at'))}. Packing list created "
                   f"{stamp(record.get('created_at'))}. TallyCraft {record.get('app_version', '')}.")
    return {
        "shop_name": (record.get("shop") or {}).get("name") or "",
        "shop_logo": logo,
        "packing_date": str(record.get("created_at", ""))[:10],
        "customer_name": cust.get("name", ""),
        "customer_address": cust.get("address", "") or "",
        "etsy_order_number": order_no,
        "order_date": cust.get("order_date", "") or "",
        "items_ordered": record["order"].get("items_ordered", ""),
        "note_to_customer": cust.get("note", "") or "",
        "footer_reference": f"Etsy order #{order_no}" if order_no else f"Packing list — {cust.get('name', '')}",
        "total_weight_lb_oz": format_lb_oz(total),
        "total_weight_g": f"{total:,.1f}",
        "unmeasured_count": skipped_n,
        "unmeasured_warning": (f"WARNING: Total weight excludes {skipped_n} part{'s' if skipped_n != 1 else ''} "
                               f"that could not be measured (marked * below). {they} in the box.")
                              if skipped_n else "",
        "parts": parts,
        "kit_headers": plan["headers"],
        "kit_legend": legend,
        "kit_totals": [str(t) for t in kit_totals],
        "total_count": str(sum(int(l.get("total_count") or 0) for l in res["lines"])),
        "total_weight_measured": f"{total:,.2f}" + (" *" if skipped_n else ""),
        "total_label": "Total (measured parts)" if skipped_n else "Total",
        "footnotes": footnotes,
        "calibration_method": "Reference piece" if cal.get("method") == "reference" else "Control sample",
        "calibration_details": details,
        "calibration_preset": cal.get("preset_name") or "",
        "grams_per_cm2": f"{g:.5f}" if isinstance(g, (int, float)) else "",
        "grams_per_in2": f"{g * 6.4516:.5f}" if isinstance(g, (int, float)) else "",
        "calculated_at": stamp(record.get("calculated_at")),
        "created_at": stamp(record.get("created_at")),
        "app_version": str(record.get("app_version", "")),
        "shop_record": "Shop record — " + record_line,
    }


# ============================================================================ render

def render_docx(record: dict, template_path, out_path, work_dir, dim_units: str = "in") -> dict:
    """Fill the template from the record and save it. Returns the kit-column plan."""
    from docx.shared import Inches
    from docxtpl import DocxTemplate, InlineImage, Listing

    work_dir = Path(work_dir)
    work_dir.mkdir(parents=True, exist_ok=True)
    geo = read_table_geometry(template_path)
    plan = plan_kit_columns(record["results"].get("kit_labels") or [], record["order"].get("packages") or [], geo)
    tpl = DocxTemplate(str(template_path))
    pic_w = (geo.before[0] if geo.found and geo.before else DEFAULT_GEOMETRY.before[0]) - 0.14

    def picture(line):
        if not (line.get("picture") or {}).get("bbox") and not (line.get("picture") or {}).get("loops"):
            return ""
        path = work_dir / f"part_{abs(hash(line['name'])) % 10**10}_{len(os.listdir(work_dir))}.png"
        render_picture_png(line.get("picture") or {}, path, pic_w)
        return InlineImage(tpl, str(path), width=Inches(pic_w))

    logo = ""
    data = decode_logo((record.get("shop") or {}).get("logo"))
    if data:
        try:
            from PIL import Image
            import io
            with Image.open(io.BytesIO(data)) as im:
                im.load()
                ext = "png" if (im.format or "").upper() == "PNG" else "jpg"
            logo_path = work_dir / f"logo.{ext}"
            logo_path.write_bytes(data)
            logo = InlineImage(tpl, str(logo_path), height=Inches(0.8))
        except Exception:
            logo = ""  # an unreadable logo is simply left out

    ctx = build_context(record, plan, picture, logo, dim_units)
    # Listing keeps line breaks; empty values stay "" so {%p if … %} hides them.
    ctx["customer_address"] = Listing(ctx["customer_address"]) if ctx["customer_address"] else ""
    ctx["note_to_customer"] = Listing(ctx["note_to_customer"]) if ctx["note_to_customer"] else ""
    tpl.render(ctx, autoescape=True)
    tpl.save(str(out_path))
    _apply_column_widths(out_path, geo, plan)
    _repair_empty_cells(out_path)
    return plan


def _repair_empty_cells(docx_path) -> None:
    """Word refuses a document ("The file appears to be corrupted") if any table
    cell has no paragraph, which happens when {%p if %} removes everything in a
    cell (e.g. no order number and no order date). Give such cells an empty one."""
    import zipfile
    from lxml import etree
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    src = Path(docx_path)
    with zipfile.ZipFile(src) as z:
        items = [(info, z.read(info.filename)) for info in z.infolist()]
    changed = False
    out = []
    for info, data in items:
        if info.filename.startswith("word/") and info.filename.endswith(".xml") and b"<w:tc" in data:
            root = etree.fromstring(data)
            for tc in root.iter(f"{{{W}}}tc"):
                if tc.find(f"{{{W}}}p") is None:
                    tc.append(etree.SubElement(tc, f"{{{W}}}p"))
                    changed = True
            data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        out.append((info, data))
    if not changed:
        return
    tmp = src.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in out:
            z.writestr(info, data)
    os.replace(tmp, src)


def _apply_column_widths(docx_path, geo: TableGeometry, plan: dict) -> None:
    """docxtpl's column loop gives every column the same width. Put the template's
    widths back and share the rest of the page among the kit columns."""
    if not geo.found:
        return
    from docx import Document
    doc = Document(str(docx_path))
    tbl = _find_parts_table(doc)
    if tbl is None:
        return
    n = len(plan["headers"])
    before = list(geo.before)
    before[geo.name_index] = max(0.8, plan["name_width"])
    widths = before + [plan["kit_width"]] * n + list(geo.after)
    twips = [round(w * 1440) for w in widths]
    _full_width_rows(tbl._tbl, len(twips))
    tables = [tbl]
    totals = _find_parts_table(doc, TOTALS_TABLE_ALT_TEXT)
    if totals is not None:
        _move_last_part_row(tbl._tbl, totals._tbl, len(twips))
        tables.append(totals)
    for t in tables:
        _set_widths(t._tbl, twips)
    doc.save(str(docx_path))


def _full_width_rows(tbl, ncols: int) -> None:
    """A row with a single cell spanning the whole table (the starter template's totals
    and sign-off row) keeps the template's span, but the column count changes with the
    number of kits: make it span every column, and drop grid columns only it needed."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for tr in tbl.findall(qn("w:tr")):
        cells = tr.findall(qn("w:tc"))
        span = cells[0].find(f"{qn('w:tcPr')}/{qn('w:gridSpan')}") if len(cells) == 1 else None
        if span is not None and int(span.get(qn("w:val"), "1")) > 1:
            span.set(qn("w:val"), str(ncols))
    grid = tbl.find(qn("w:tblGrid"))
    cols = grid.findall(qn("w:gridCol"))
    widest = max((sum(int(tc.find(f"{qn('w:tcPr')}/{qn('w:gridSpan')}").get(qn("w:val")))
                      if tc.find(f"{qn('w:tcPr')}/{qn('w:gridSpan')}") is not None else 1
                      for tc in tr.findall(qn("w:tc"))) for tr in tbl.findall(qn("w:tr"))), default=0)
    if widest == ncols:
        for col in cols[ncols:]:
            grid.remove(col)
        for _ in range(ncols - len(cols)):
            grid.append(OxmlElement("w:gridCol"))


def _move_last_part_row(parts, totals, ncols: int) -> None:
    """Move the last part row into the totals table (which sits in the parts table's
    unsplittable last row), so at least one part always goes to the next page with
    the totals and sign-off instead of them starting a page alone."""
    from docx.oxml.ns import qn
    end_row = next((tr for tr in parts.findall(qn("w:tr")) if totals in tr.iter(qn("w:tbl"))), None)
    prev = end_row.getprevious() if end_row is not None else None
    first = totals.find(qn("w:tr"))
    if (prev is None or prev.tag != qn("w:tr") or first is None or len(prev.findall(qn("w:tc"))) != ncols
            or prev.find(f"{qn('w:trPr')}/{qn('w:tblHeader')}") is not None):
        return
    first.addprevious(prev)


def _set_widths(tbl, twips: list[int]) -> None:
    """Grid, cell, and table widths for one table (skipped if it has a different number of columns)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    grid = tbl.find(qn("w:tblGrid"))
    cols = grid.findall(qn("w:gridCol"))
    if len(cols) != len(twips):
        return  # the table was restructured; leave docxtpl's layout alone
    for col, w in zip(cols, twips):
        col.set(qn("w:w"), str(w))
    for tr in tbl.findall(qn("w:tr")):
        cells = tr.findall(qn("w:tc"))
        if len(cells) == len(twips):
            ws = twips
        elif len(cells) == 1:
            ws = [sum(twips)]  # a full-width row
        else:
            continue
        for tc, w in zip(cells, ws):
            pr = tc.find(qn("w:tcPr"))
            if pr is None:
                pr = OxmlElement("w:tcPr")
                tc.insert(0, pr)
            tcw = pr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                pr.insert(0, tcw)
            tcw.set(qn("w:w"), str(w))
            tcw.set(qn("w:type"), "dxa")
    pr = tbl.find(qn("w:tblPr"))
    tblw = pr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        pr.append(tblw)
    tblw.set(qn("w:w"), str(sum(twips)))
    tblw.set(qn("w:type"), "dxa")


# ============================================================================ validation

def validate_template(path) -> tuple[list[str], list[str]]:
    """(errors, warnings). Errors block generation; warnings (likely typos) can be skipped."""
    from docxtpl import DocxTemplate
    from jinja2 import TemplateSyntaxError

    errors: list[str] = []
    warnings: list[str] = []
    p = Path(path)
    if not p.exists():
        return [f"The packing list template wasn't found: {p}"], []
    try:
        tpl = DocxTemplate(str(p))
        tpl.init_docx()
    except Exception as exc:
        return [f"The template {p.name} couldn't be opened as a Word document ({exc})."], []
    try:
        used = set(tpl.get_undeclared_template_variables())
    except TemplateSyntaxError as exc:
        return [f"The template {p.name} has a broken placeholder: {exc.message}. Check the brackets "
                "{{ }} and {% %} near it."], []
    except Exception as exc:
        return [f"The template {p.name} couldn't be read ({exc})."], []

    if "parts" not in used:
        errors.append("The parts table loop is missing: add a row containing {%tr for p in parts %} "
                      "(and a row with {%tr endfor %}).")
    if not used & {"total_weight_lb_oz", "total_weight_g"}:
        errors.append("The total weight is missing: add {{ total_weight_lb_oz }} or {{ total_weight_g }}.")
    for name in sorted(used - KNOWN_FIELDS):
        warnings.append(f"Unrecognized placeholder {{{{ {name} }}}} — possibly a typo. It will be left blank.")

    # Loop-item attributes (p.name, f.reason) aren't checked by docxtpl; check them here.
    text = _template_text(tpl)
    for loop_var, collection in re.findall(r"\{%-?\s*(?:tr|tc|p)?\s*for\s+(\w+)\s+in\s+(\w+)\s*-?%\}", text):
        allowed = PART_ATTRS if collection == "parts" else FOOTNOTE_ATTRS if collection == "footnotes" else None
        if allowed is None:
            continue
        for attr in sorted(set(re.findall(rf"\b{loop_var}\.(\w+)", text)) - allowed):
            warnings.append(f"Unrecognized placeholder {{{{ {loop_var}.{attr} }}}} — possibly a typo. "
                            "It will be left blank.")
    return errors, warnings


def _template_text(tpl) -> str:
    """All Jinja text in the document, headers, and footers (after docxtpl's tag clean-up)."""
    parts = [tpl.get_xml()]
    try:
        for rel in tpl.docx.part.rels.values():
            if rel.reltype.endswith("/header") or rel.reltype.endswith("/footer"):
                parts.append(rel.target_part.blob.decode("utf-8", "ignore"))
    except Exception:
        pass
    joined = "".join(parts)
    try:
        joined = tpl.patch_xml(joined)
    except Exception:
        pass
    return re.sub(r"<[^>]+>", "", joined).replace("&quot;", '"').replace("&amp;", "&")
