"""Packing lists from the Word template (docxtpl). Word itself is never needed:
PDF conversion is mocked, and the filled-in .docx is inspected directly."""

import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from tallycraft.order import Order, Package, merge_order
from tallycraft.packing_docx import (DEFAULT_GEOMETRY, FIELDS, KNOWN_FIELDS, bundled_template, build_context,
                                     field_reference_markdown, field_reference_text, plan_kit_columns,
                                     read_table_geometry, render_docx, render_picture_png, validate_template)
from tallycraft.pieces import PieceRow
from tallycraft.storage import PresetError, Storage
from test_packing import add, make_record  # noqa: E402
from test_packing import order  # noqa: F401,F811  (pytest fixture)

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = bundled_template()


def docx_text(path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:  # every paragraph, including tables inside cells (the totals/sign-off row)
        parts += ["".join(x.text or "" for x in p.iter(qn("w:t"))) for p in t._tbl.iter(qn("w:p"))]
    parts += [p.text for p in doc.sections[0].footer.paragraphs]
    return "\n".join(parts)


def parts_table(path):
    doc = Document(str(path))
    return next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Picture")


def table_rows(path):
    """Header, part rows and totals row in order, including the last part row and the
    totals, which sit in a table inside the unsplittable end row (with the sign-off)."""
    t = parts_table(path)
    nested = t.rows[-1].cells[0].tables
    return list(t.rows[:-1]) + (list(nested[0].rows) if nested else [t.rows[-1]])


def fill(record, tmp_path, name="out.docx", template=TEMPLATE):
    out = tmp_path / name
    plan = render_docx(record, template, out, tmp_path / f"{name}_work")
    return out, plan


# ---------------------------------------------------------------- filling in

def test_bundled_template_is_valid():
    assert TEMPLATE.exists()
    assert validate_template(TEMPLATE) == ([], [])


def test_fill_multi_kit_order(order, tmp_path):
    rec = make_record(order)
    out, plan = fill(rec, tmp_path)
    text = docx_text(out)
    for expected in ("Jordan Rivera", "3141592", "2026-09-27", "2x XL Standard Box, 1x XL Tunnel + Ramp",
                     "Bun Homes Co.", "Thanks!", "Etsy order #3141592", "Shop record — Calibration: control sample"):
        assert expected in text, expected
    assert plan["mode"] == "short"  # the starter template's kit column is too narrow for full names
    assert "Kit columns: Kit 1 = XL Standard Box (1 of 2); Kit 2 = XL Standard Box (2 of 2)" in text
    t = parts_table(out)
    header = [c.text for c in t.rows[0].cells]
    assert header[:2] == ["Picture", "Part"]
    assert header[2:5] == ["Kit 1", "Kit 2", "Kit 3"]
    assert header[5:] == ["Total count", "Weight per piece (g)", "Total weight (g)", "Done"]
    # header + one row per part (side, floor, gap) + totals row
    rows = table_rows(out)
    assert len(rows) == 1 + 3 + 1
    side = next(r for r in rows if r.cells[1].text.startswith("side"))
    assert [c.text for c in side.cells[2:6]] == ["2", "2", "3", "7"]
    floor = next(r for r in rows if r.cells[1].text.startswith("floor"))
    assert [c.text for c in floor.cells[2:5]] == ["1", "1", "—"]


def test_single_kit_order(order, tmp_path):
    single = Order([Package("XL Standard Box", preset_name="XL Standard Box")])
    for iid in order.packages[0].order[:2]:
        add(single.packages[0], order.packages[0].rows[iid])
    out, plan = fill(make_record(single), tmp_path)
    assert plan["mode"] == "single"
    header = [c.text for c in parts_table(out).rows[0].cells]
    # one kit: no "Kit 1" column next to a "Total count" column, just "Count"
    assert header == ["Picture", "Part", "Count", "Weight per piece (g)", "Total weight (g)", "Done"]
    rows = table_rows(out)
    assert all(len(r.cells) == 6 for r in rows)
    assert rows[1].cells[2].text == "2"
    text = docx_text(out)
    assert "WARNING" not in text  # {%p if unmeasured_warning %} removed the paragraph
    assert "Kit columns" not in text


def test_column_widths_follow_the_template(order, tmp_path):
    out, plan = fill(make_record(order), tmp_path)
    t = parts_table(out)
    grid = [int(g.get(qn("w:w"))) / 1440 for g in t._tbl.tblGrid.findall(qn("w:gridCol"))]
    geo = read_table_geometry(TEMPLATE)
    assert grid[0] == pytest.approx(geo.before[0], abs=0.01)  # picture column kept
    assert grid[-4:] == pytest.approx(geo.after, abs=0.01)  # count/weight/checkbox kept
    assert sum(grid) == pytest.approx(geo.content_width, abs=0.02)  # fills the page, never wider
    assert grid[2] == grid[3] == grid[4]  # kit columns share the rest equally


@pytest.mark.parametrize("qty_a, qty_b", [(1, 0), (2, 1), (4, 4), (15, 0), (40, 40)])
def test_table_end_stays_with_the_sign_off(order, tmp_path, qty_a, qty_b):
    """The last part row, totals, footnotes, sign-off, note and shop record share one
    unsplittable row that spans every column, whatever the number of kit columns."""
    order.packages[0].quantity = qty_a
    if qty_b:
        order.packages[1].quantity = qty_b
    else:
        del order.packages[1]
    out, plan = fill(make_record(order), tmp_path)
    t = parts_table(out)
    ncols = len(t.rows[0].cells)
    grid = t._tbl.tblGrid.findall(qn("w:gridCol"))
    assert len(grid) == ncols
    end = t.rows[-1]._tr
    assert end.find(f"{qn('w:trPr')}/{qn('w:cantSplit')}") is not None
    tcs = end.findall(qn("w:tc"))
    assert len(tcs) == 1 and int(tcs[0].find(f"{qn('w:tcPr')}/{qn('w:gridSpan')}").get(qn("w:val"))) == ncols
    cell = t.rows[-1].cells[0]
    inner = cell.tables[0]
    assert len(inner.rows) == 2  # the last part row + the totals row
    assert inner.rows[1].cells[1].text.startswith("Total")
    assert [g.get(qn("w:w")) for g in inner._tbl.tblGrid.findall(qn("w:gridCol"))] == [g.get(qn("w:w")) for g in grid]
    text = "\n".join(p.text for p in cell.paragraphs)
    assert "Packed by:" in text and "Shop record" in text
    assert "Thanks!" in cell.tables[1].rows[0].cells[0].text  # the note box
    assert len(t.rows) == 1 + (len(table_rows(out)) - 1 - 2) + 1  # header, other parts, end row


def test_unmeasured_parts(order, tmp_path):
    out, _plan = fill(make_record(order), tmp_path)
    text = docx_text(out)
    assert "WARNING: Total weight excludes 1 part that could not be measured" in text
    assert "*1 gap: could not be measured" in text
    gap = next(r for r in table_rows(out) if r.cells[1].text.startswith("gap"))
    assert gap.cells[1].text.endswith("*1")
    assert [c.text for c in gap.cells[6:8]] == ["—", "—"]
    assert "Total (measured parts)" in text


@pytest.mark.parametrize("qty_a, qty_b, mode, first_header", [
    (4, 4, "short", "Kit 1"),
    (15, 0, "per_package", "Kit 1"),  # "Kit 1 = XL Standard Box (each of 15)" in the legend
])
def test_many_kits(order, tmp_path, qty_a, qty_b, mode, first_header):
    order.packages[0].quantity = qty_a
    if qty_b:
        order.packages[1].quantity = qty_b
    else:
        order.packages.pop()
    out, plan = fill(make_record(order), tmp_path)
    assert plan["mode"] == mode
    t = parts_table(out)
    assert t.rows[0].cells[2].text == first_header
    grid = [int(g.get(qn("w:w"))) / 1440 for g in t._tbl.tblGrid.findall(qn("w:gridCol"))]
    assert sum(grid) <= 7.5 + 0.02
    text = docx_text(out)
    if mode == "short":
        assert "Kit columns: Kit 1 = XL Standard Box (1 of 4)" in text
        assert len(t.rows[0].cells) == 2 + 8 + 4
    else:
        assert "counts are shown once per package" in text
        assert len(t.rows[0].cells) == 2 + 1 + 4


def test_optional_fields_left_out(order, tmp_path):
    rec = make_record(order, customer={"name": "Sam", "address": "", "etsy_order": "", "order_date": "", "note": ""},
                      shop={"name": "", "logo": None})
    text = docx_text(fill(rec, tmp_path)[0])
    assert "ETSY ORDER #" not in text and "ORDER DATE" not in text
    assert "Packing list — Sam" in text


def test_no_empty_table_cells_word_would_reject(order, tmp_path):
    """Word calls a document corrupted if a table cell has no paragraph; {%p if %} can
    empty a cell when optional fields are blank. Checked in body and footer XML."""
    from lxml import etree
    rec = make_record(order, customer={"name": "Sam", "address": "", "etsy_order": "", "order_date": "", "note": ""},
                      shop={"name": "", "logo": None})
    out, _ = fill(rec, tmp_path)
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(out) as z:
        for name in z.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                root = etree.fromstring(z.read(name))
                empty = [tc for tc in root.iter(f"{W}tc") if tc.find(f"{W}p") is None]
                assert not empty, f"{len(empty)} empty cell(s) in {name}"
    doc = Document(str(out))
    ship = next(c for t in doc.tables for c in t.rows[0].cells if c.paragraphs and c.paragraphs[0].text == "SHIP TO")
    assert [p.text for p in ship.paragraphs] == ["SHIP TO", "Sam"]  # no blank address line either


def test_empty_cells_are_repaired_for_edited_templates(order, tmp_path):
    """A user's template where a whole cell is conditional still produces a valid document."""
    def conditional_cell(d):
        t = d.tables[0]  # header block: put the date cell behind a condition that's false
        cell = t.rows[0].cells[1]
        for p in list(cell.paragraphs):
            p._p.getparent().remove(p._p)
        cell.add_paragraph("{%p if etsy_order_number %}")
        cell.add_paragraph("{{ etsy_order_number }}")
        cell.add_paragraph("{%p endif %}")
    custom = _template_with(tmp_path, "cond.docx", conditional_cell)
    rec = make_record(order, customer={"name": "Sam", "address": "", "etsy_order": "", "order_date": "", "note": ""})
    out, _ = fill(rec, tmp_path, template=custom)
    xml = zipfile.ZipFile(out).read("word/document.xml")
    assert b"<w:tc><w:tcPr" not in xml or b"</w:tcPr></w:tc>" not in xml


def test_special_characters_and_address_line_breaks(order, tmp_path):
    rec = make_record(order, customer={"name": "Jo & Sam <Rivera>", "address": "12 Oak St\nApt 3\nSpringfield",
                                       "etsy_order": "7", "order_date": "2026-09-27", "note": "A < B & C"})
    out, _ = fill(rec, tmp_path)
    text = docx_text(out)
    assert "Jo & Sam <Rivera>" in text and "A < B & C" in text
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    addr = xml[xml.index("12 Oak St"):xml.index("Springfield")]
    assert addr.count("<w:br/>") == 2  # line breaks kept


def test_pictures_and_logo_embedded(order, tmp_path):
    png = (ROOT / "assets" / "tallycraft.png").read_bytes()
    from tallycraft.packing import encode_logo
    rec = make_record(order, shop={"name": "Shop", "logo": encode_logo(png, "logo.png")})
    out, _ = fill(rec, tmp_path)
    media = [n for n in zipfile.ZipFile(out).namelist() if n.startswith("word/media/")]
    assert len(media) >= 1 + 3  # logo + a picture per part


def test_unknown_placeholders_render_blank(tmp_path, order):
    d = Document(str(TEMPLATE))
    d.add_paragraph("Extra: [{{ not_a_field }}]")
    custom = tmp_path / "custom.docx"
    d.save(custom)
    out, _ = fill(make_record(order), tmp_path, template=custom)
    assert "Extra: []" in docx_text(out)


# ---------------------------------------------------------------- pictures

def test_picture_png_resolution_and_colors(tmp_path):
    pic = {"loops": [[[0, 0], [10, 0], [10, 10], [0, 10]], [[4, 4], [6, 4], [6, 6], [4, 6]]], "depths": [0, 1],
           "open": [], "engrave": [[[1, 1], [9, 9]]], "bbox": [10, 10]}
    path = tmp_path / "p.png"
    render_picture_png(pic, path, width_in=0.8, height_in=0.72)
    with Image.open(path) as im:
        assert im.size == (round(0.8 * 400), round(0.72 * 400))  # 400 DPI at printed size
        assert im.info["dpi"][0] >= 300
        colors = {c for _n, c in im.convert("RGB").getcolors(1_000_000)}
    assert (0, 0, 0) in colors  # cut lines solid black
    assert any(120 < r < 200 and r == g == b for r, g, b in colors)  # engrave lighter grey


def test_picture_without_geometry(tmp_path):
    path = tmp_path / "empty.png"
    render_picture_png({"loops": [], "open": [], "engrave": [], "bbox": None}, path, 0.8)
    assert path.exists()


# ---------------------------------------------------------------- kit plan

@pytest.mark.parametrize("n, packages, mode", [
    (3, None, "full"), (8, None, "short"),
    (15, [{"name": "XL Standard Box", "quantity": 15}], "per_package"),
    (40, [{"name": f"P{i}", "quantity": 2} for i in range(20)], "omitted"),
])
def test_plan_modes(n, packages, mode):
    labels = [f"XL Standard Box ({i} of {n})" for i in range(1, n + 1)]
    plan = plan_kit_columns(labels, packages, DEFAULT_GEOMETRY)
    assert plan["mode"] == mode
    assert plan["name_width"] + plan["kit_width"] * len(plan["headers"]) <= DEFAULT_GEOMETRY.content_width - 0 + 1e-9 \
        or mode == "omitted"


# ---------------------------------------------------------------- validation

def _template_with(tmp_path, name, transform):
    d = Document(str(TEMPLATE))
    transform(d)
    path = tmp_path / name
    d.save(path)
    return path


def _replace_everywhere(d, old, new):
    for p in list(d.paragraphs) + [p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs]:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new)


def test_validation_missing_parts_loop(tmp_path):
    def drop_table(d):
        t = next(t for t in d.tables if t.rows[0].cells[0].text.strip() == "Picture")
        t._tbl.getparent().remove(t._tbl)
    errors, _ = validate_template(_template_with(tmp_path, "noparts.docx", drop_table))
    assert any("parts table loop is missing" in e for e in errors)


def test_validation_missing_total_weight(tmp_path):
    def drop_total(d):
        _replace_everywhere(d, "{{ total_weight_lb_oz }}", "")
        _replace_everywhere(d, "({{ total_weight_g }} g)", "")
    errors, _ = validate_template(_template_with(tmp_path, "nototal.docx", drop_total))
    assert any("total weight is missing" in e for e in errors)


def test_validation_typos_are_warnings(tmp_path):
    def typos(d):
        _replace_everywhere(d, "{{ customer_name }}", "{{ custmer_name }}")
        _replace_everywhere(d, "{{ p.total_count }}", "{{ p.totl_count }}")
    errors, warnings = validate_template(_template_with(tmp_path, "typo.docx", typos))
    assert errors == []
    assert any("{{ custmer_name }}" in w for w in warnings)
    assert any("{{ p.totl_count }}" in w for w in warnings)


def test_validation_broken_brackets(tmp_path):
    errors, _ = validate_template(_template_with(
        tmp_path, "broken.docx", lambda d: d.add_paragraph("{{ customer_name ")))
    assert errors and "broken placeholder" in errors[0]


def test_validation_unopenable(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a zip")
    assert "couldn't be opened" in validate_template(bad)[0][0]
    assert "wasn't found" in validate_template(tmp_path / "missing.docx")[0][0]


# ---------------------------------------------------------------- the field reference

def test_every_field_is_filled_and_documented(order):
    ctx = build_context(make_record(order), plan_kit_columns([], None, DEFAULT_GEOMETRY))
    assert set(ctx) == KNOWN_FIELDS


def test_docs_are_generated_from_fields():
    assert (ROOT / "docs" / "TEMPLATE_FIELDS.md").read_text(encoding="utf-8").strip() == \
        field_reference_markdown().strip()
    readme = (ROOT / "dist_template" / "README.txt").read_text(encoding="utf-8")
    assert field_reference_text() in readme
    for f in FIELDS:
        assert f"{{{{ {f.name} }}}}" in readme


# ---------------------------------------------------------------- Word missing

def test_word_missing_gives_plain_message(monkeypatch, tmp_path):
    from tallycraft import word_pdf
    monkeypatch.setitem(sys.modules, "win32com", None)
    monkeypatch.setitem(sys.modules, "win32com.client", None)
    with pytest.raises(word_pdf.PdfConversionError, match="Word automation isn't available"):
        word_pdf.convert_docx_to_pdf(tmp_path / "a.docx", tmp_path / "a.pdf")


# ---------------------------------------------------------------- template storage

def test_templates_never_overwrite_user_edits(tmp_path):
    s = Storage(tmp_path)
    s.ensure_templates(TEMPLATE)
    user = s.templates_dir / "packing_list_template.docx"
    pristine = s.templates_dir / "_default" / "packing_list_template.docx"
    assert user.read_bytes() == pristine.read_bytes() == TEMPLATE.read_bytes()
    user.write_bytes(b"my edited template")
    s.ensure_templates(TEMPLATE)  # e.g. after an update
    assert user.read_bytes() == b"my edited template"
    pristine.write_bytes(b"stale")
    s.ensure_templates(TEMPLATE)
    assert pristine.read_bytes() == TEMPLATE.read_bytes()  # the app's own copy is refreshed


def test_restore_default_makes_new_copies(tmp_path):
    s = Storage(tmp_path)
    s.ensure_templates(TEMPLATE)
    a = s.restore_default_template()
    b = s.restore_default_template()
    assert a.name == "packing_list_template (default copy).docx"
    assert b.name == "packing_list_template (default copy) 2.docx"
    assert a.read_bytes() == TEMPLATE.read_bytes()


def test_restore_without_default_is_a_friendly_error(tmp_path):
    with pytest.raises(PresetError, match="default template is missing"):
        Storage(tmp_path).restore_default_template()


def test_template_path_and_reprint_names(tmp_path):
    s = Storage(tmp_path)
    assert s.template_path({"template_path": "templates/x.docx"}) == tmp_path / "templates" / "x.docx"
    assert s.template_path({"template_path": str(tmp_path / "abs.docx")}) == tmp_path / "abs.docx"
    s.ensure_folders()
    j = s.packing_dir / "2026-09-27 - Sam - 1.json"
    d1, p1 = s.reprint_paths(j, "2026-09-28 1015")
    assert d1.name == "2026-09-27 - Sam - 1 (reprint 2026-09-28 1015).docx" and p1.suffix == ".pdf"
    d1.write_bytes(b"x")
    d2, _ = s.reprint_paths(j, "2026-09-28 1015")
    assert d2.name == "2026-09-27 - Sam - 1 (reprint 2026-09-28 1015 2).docx"


def test_packing_names_avoid_existing_docx(tmp_path):
    s = Storage(tmp_path)
    s.ensure_folders()
    (s.packing_dir / "A.docx").write_bytes(b"x")
    json_path, _ = s.packing_paths("A")
    assert json_path.name == "A (2).json"


def test_lightburn_panel_picture_has_engrave(tmp_path):
    lb = ROOT / "tests" / "fixtures_private" / "Main Box - Front Panel - LB Test.dxf"
    if not lb.exists():
        pytest.skip("private LightBurn sample not present")
    from tallycraft.packing import picture_from_parsed
    pic = picture_from_parsed(PieceRow.load(str(lb)).parsed)
    assert len(pic["engrave"]) == 151 and len(pic["loops"]) == 23
    render_picture_png(pic, tmp_path / "lb.png", 0.81)


# ---------------------------------------------------------------- display names & dimension units (v0.5)

def test_display_name_on_packing_list(order, tmp_path):
    side = next(r for _p, _i, r in order.all_rows() if r.name == "side.dxf")
    for _p, _i, r in order.all_rows():
        if r.name == "side.dxf":
            r.display_name = "Side Wall"
        if r.name == "gap.dxf":
            r.display_name = "Broken Bit"
    _labels, merged = merge_order(order)
    rec = make_record(order)
    for line in rec["results"]["lines"]:  # make_record builds pieces_info without display names
        line["display_name"] = next(m.display_name for m in merged if m.name == line["name"])
    ctx = build_context(rec, {"mode": "full", "headers": [], "indices": [], "legend": []})
    by_file = {p["name"]: p for p in ctx["parts"]}
    assert by_file["side"]["display_name"] == "Side Wall"  # p.display_name: the name to show
    assert by_file["floor"]["display_name"] == "floor"  # no display name: the file name
    assert ctx["footnotes"][0]["name"] == "Broken Bit"
    out, _plan = fill(rec, tmp_path)
    text = docx_text(out)
    assert "Side Wall" in text and "Broken Bit *1" in text and side.name not in text


@pytest.mark.parametrize("unit, show, expected", [
    ("in", "in", "4 × 2 in"),
    ("in", "mm", "102 × 50.8 mm"),
    ("mm", "in", "3.94 × 1.97 in"),
    ("mm", "mm", "100 × 50 mm"),
    ("mm", "file", "100 × 50 mm"),
    ("in", "file", "4 × 2 in"),
    (None, "in", "4 × 2"),  # unknown units can't be converted
])
def test_dimension_units(unit, show, expected):
    from tallycraft.packing_docx import _dims
    bbox = [4.0, 2.0] if unit != "mm" else [100.0, 50.0]
    assert _dims({"unit": unit, "picture": {"bbox": bbox}}, show) == expected


def test_dimension_units_setting_reaches_the_document(order, tmp_path):
    rec = make_record(order)
    out = tmp_path / "mm.docx"
    render_docx(rec, TEMPLATE, out, tmp_path / "work", "mm")
    assert "102 × 50.8 mm" in docx_text(out)


def test_dimension_units_setting_is_sanitized(tmp_path):
    s = Storage(tmp_path)
    s.ensure_folders()
    s.settings_path.write_text('{"dimension_units": "furlongs"}', encoding="utf-8")
    assert s.load_settings()[0]["dimension_units"] == "in"


# ---------------------------------------------------------------- v0.6: numbers, single kit, pictures

@pytest.mark.parametrize("value, text", [
    (280.19, "280"), (6.27, "6.27"), (0.78, "0.78"), (20.13, "20.1"), (204.85, "205"), (268.40, "268"),
    (14.5, "14.5"), (14.25, "14.3"), (0.734, "0.73"), (0.735, "0.74"), (13.661, "13.7"), (8.125, "8.13"),
    (9.996, "10"), (999.6, "1,000"), (1694.58, "1,695"), (12564.18, "12,564"), (0, "0"), (0.004, "<0.01"),
])
def test_short_number(value, text):
    from tallycraft.packing_docx import short_number
    assert short_number(value) == text


def test_numbers_on_the_packing_list(order):
    rec = make_record(order)
    rec["results"]["total_g"] = 5149.2
    plan = {"mode": "full", "headers": [], "indices": [], "legend": []}
    ctx = build_context(rec, plan)
    assert ctx["total_weight_lb_oz"] == "11 lb 6 oz"  # 11 lb 5.6 oz: whole ounces, rounded up
    assert ctx["total_weight_measured"] == "5,149 *"  # one part is unmeasured
    assert ctx["count_header"] == "Total count"
    side = next(p for p in ctx["parts"] if p["name"] == "side")
    line = next(l for l in rec["results"]["lines"] if l["name"] == "side.dxf")
    from tallycraft.packing_docx import short_number
    assert side["weight_per_piece"] == short_number(line["weight_per_piece_g"])
    assert side["total_weight"] == short_number(line["item_total_g"])
    assert build_context(rec, {**plan, "mode": "single"})["count_header"] == "Count"


def test_single_kit_plan():
    plan = plan_kit_columns(["XL Standard Box"], [{"name": "XL Standard Box", "quantity": 1}], DEFAULT_GEOMETRY)
    assert plan["mode"] == "single" and plan["headers"] == [] and plan["indices"] == [] and plan["legend"] == []
    assert plan_kit_columns(["A (1 of 2)", "A (2 of 2)"], None, DEFAULT_GEOMETRY)["mode"] != "single"


def test_starter_template_fields():
    from docxtpl import DocxTemplate
    from tallycraft.packing_docx import _template_text
    tpl = DocxTemplate(str(TEMPLATE))
    tpl.init_docx()
    text = _template_text(tpl)
    assert "{{ p.display_name }}" in text and "{{ p.name }}" not in text  # Display name when set
    assert "{{ count_header }}" in text
    assert "total_weight_g" not in text  # pounds and ounces only


def test_pictures_fit_inside_their_cell_and_are_centred(order, tmp_path):
    out, _plan = fill(make_record(order), tmp_path)
    geo = read_table_geometry(TEMPLATE)
    room = geo.before[0] - geo.picture_margins
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    import re
    inlines = re.findall(r"<wp:inline[^>]*>.*?<wp:extent cx=\"(\d+)\"", xml)
    assert inlines
    for cx in inlines:
        assert int(cx) / 914400 <= room  # wider than the cell's text area: Word pushes it off centre
    for tag in re.findall(r"<wp:inline[^>]*>", xml):  # LibreOffice adds left spacing unless it's 0
        assert all(f'{side}="0"' in tag for side in ("distT", "distB", "distL", "distR"))


def test_picture_margins_come_from_the_table_style():
    from tallycraft.packing_docx import _picture_margins, _find_parts_table
    doc = Document(str(TEMPLATE))
    assert _picture_margins(doc, _find_parts_table(doc)) == pytest.approx(2 * 108 / 1440)  # Word's default


def test_picture_background_is_transparent(tmp_path):
    path = tmp_path / "p.png"
    render_picture_png({"loops": [[[0, 0], [10, 0], [10, 10], [0, 10]]], "depths": [0], "open": [], "engrave": [],
                        "bbox": [10, 10]}, path, width_in=0.8)
    with Image.open(path) as im:
        assert im.mode == "RGBA"
        assert im.getpixel((0, 0))[3] == 0  # striped rows show through
        assert im.getchannel("A").getextrema()[1] == 255  # the outline is solid


@pytest.mark.parametrize("nparts, shaded", [(3, True), (2, False)])
def test_moved_part_row_keeps_its_stripe(order, tmp_path, nparts, shaded):
    """The last part row moves into the totals table; the parts table's row stripes
    (1st, 3rd, … part shaded) must carry on there."""
    from tallycraft.packing_docx import _band_fill
    rec = make_record(order)
    rec["results"]["lines"] = rec["results"]["lines"][:nparts]
    out, _plan = fill(rec, tmp_path)
    doc = Document(str(out))
    t = parts_table(out)
    fill_color = _band_fill(doc, t._tbl)
    assert fill_color  # the starter template stripes its rows
    moved = t.rows[-1].cells[0].tables[0].rows[0]._tr
    fills = {tc.find(f"{qn('w:tcPr')}/{qn('w:shd')}").get(qn("w:fill"))
             if tc.find(f"{qn('w:tcPr')}/{qn('w:shd')}") is not None else None for tc in moved.findall(qn("w:tc"))}
    assert fills == ({fill_color} if shaded else {None})

